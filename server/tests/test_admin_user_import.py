"""T15 批量导入开户测试：模板下载 / 后台线程逐行校验（格式/查重/区域/角色/scope）/ 台账轮询 / 初始密码后6位登录
口径：计划 T15 Step 1（4 行 xlsx：1 正常 / 1 手机号重复 / 1 区域码不存在 / 1 角色非法 → success==1 fail==3、
detail 逐行 msg 准确、轮询返回台账、导入用户可用后 6 位密码登录）+ 模板/格式/批内重复/scope/空表/坏文件补充。
后台线程说明：实现用 threading.Thread(daemon=True)（分发手册规定），POST 返回后线程异步跑；
线程会话走模块级 SessionLocal，测试 monkeypatch 到 db fixture 同一引擎（tests/ 无 __init__.py，
import tests.conftest 会得到与 pytest 所加载不同的模块实例，故从 db.get_bind() 现取引擎——P-export 经验）。
"""
import io
import time

import pytest
from openpyxl import Workbook, load_workbook

from app.core.security import create_token
from app.models import Region, User, UserImportBatch
from tests.conftest import make_user


# ---------- 造数 ----------

def seed_regions(db):
    db.add_all([
        Region(code="330000", name="浙江省", level="province", parent_code=None),
        Region(code="331000", name="台州市", level="city", parent_code="330000"),
        Region(code="331004", name="临海市", level="district", parent_code="331000"),
        Region(code="331082", name="三门县", level="district", parent_code="331000"),
        Region(code="330100", name="杭州市", level="city", parent_code="330000"),
        Region(code="330105", name="拱墅区", level="district", parent_code="330100"),
    ])
    db.commit()


def auth_of(user):
    return {"Authorization": f"Bearer {create_token(str(user.id))}"}


@pytest.fixture()
def import_env(db, monkeypatch):
    """后台线程会话 → 与 db fixture 同一引擎（db.get_bind()）"""
    from sqlalchemy.orm import sessionmaker

    from app.api.admin import user_import as ui
    monkeypatch.setattr(ui, "SessionLocal",
                        sessionmaker(bind=db.get_bind(), expire_on_commit=False))
    return ui


def xlsx_bytes(rows, headers=("手机号", "姓名", "区域码", "单位", "角色")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def wait_batch(client, headers, batch_id, timeout=5.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        r = client.get(f"/api/admin/users/import/{batch_id}", headers=headers)
        assert r.status_code == 200
        d = r.json()["data"]
        if d["status"] == "completed":
            return d
        time.sleep(0.05)
    pytest.fail("导入批次超时未完成")


# ---------- 模板 ----------

def test_import_template(client, db):
    seed_regions(db)
    admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    r = client.get("/api/admin/users/import-template", headers=auth_of(admin))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert len(r.content) > 0
    wb = load_workbook(io.BytesIO(r.content))
    assert [c.value for c in wb.active[1]] == ["手机号", "姓名", "区域码", "单位", "角色"]


# ---------- 核心：4 行混合校验 ----------

def test_import_mixed_rows_and_login(client, db, import_env):
    seed_regions(db)
    make_user(db, phone="33100400002", name="已有民警")            # 供库内查重
    admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    h = auth_of(admin)

    rows = [
        ("33100400099", "王五", "331004", "临海市公安局", "user"),    # 正常
        ("33100400002", "重复号", "331004", "临海市公安局", "user"),  # 库内重复
        ("33100400098", "赵六", "999999", "临海市公安局", "user"),    # 区域不存在
        ("33100400097", "钱七", "331004", "临海市公安局", "boss"),    # 角色非法
    ]
    r = client.post("/api/admin/users/import", headers=h, files={
        "file": ("users.xlsx", xlsx_bytes(rows),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200 and r.json()["code"] == 0
    batch_id = r.json()["data"]["batch_id"]

    d = wait_batch(client, h, batch_id)
    assert d["total"] == 4 and d["success"] == 1 and d["fail"] == 3
    detail = {row["phone"]: row for row in d["detail"]}
    assert detail["33100400099"]["ok"] is True
    assert "已存在" in detail["33100400002"]["msg"]
    assert "区域" in detail["33100400098"]["msg"]
    assert "角色" in detail["33100400097"]["msg"]

    # 台账落库
    batch = db.get(UserImportBatch, batch_id)
    assert batch is not None and batch.success == 1 and batch.fail == 3 and batch.total == 4

    # 成功用户可用手机号后 6 位登录，且带 import_batch_id
    login = client.post("/api/auth/login", json={"phone": "33100400099", "password": "400099"})
    assert login.status_code == 200 and login.json()["data"]["token"]
    u = db.query(User).filter_by(phone="33100400099").first()
    assert u is not None and u.import_batch_id == batch_id


# ---------- 格式 / 批内重复 / 姓名空 ----------

def test_import_phone_format_in_batch_dup_and_empty_name(client, db, import_env):
    seed_regions(db)
    admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    h = auth_of(admin)

    rows = [
        ("12345", "短号", "331004", "", "user"),            # 手机号格式
        ("33100400011", "李八", "331004", "", "user"),      # 首次成功
        ("33100400011", "李八2", "331004", "", "user"),     # 批内重复
        ("33100400012", "", "331004", "", "user"),          # 姓名为空
    ]
    batch_id = client.post("/api/admin/users/import", headers=h, files={
        "file": ("users.xlsx", xlsx_bytes(rows),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }).json()["data"]["batch_id"]

    d = wait_batch(client, h, batch_id)
    assert d["total"] == 4 and d["success"] == 1 and d["fail"] == 3
    detail = {row["phone"]: row["msg"] for row in d["detail"]}
    assert "11位" in detail["12345"]
    assert "已存在" in detail["33100400011"]
    assert "姓名" in detail["33100400012"]


# ---------- scope：admin 仅可导本辖区 ----------

def test_import_scope_restriction(client, db, import_env):
    seed_regions(db)
    county_admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    super_admin = make_user(db, phone="33000000001", name="省超管", role="super_admin", region="330000")

    # 县管导外辖区 → 该行失败
    batch_id = client.post("/api/admin/users/import", headers=auth_of(county_admin), files={
        "file": ("users.xlsx", xlsx_bytes([("33108200077", "越界", "331082", "", "user")]),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }).json()["data"]["batch_id"]
    d = wait_batch(client, auth_of(county_admin), batch_id)
    assert d["success"] == 0 and d["fail"] == 1
    assert "范围" in d["detail"][0]["msg"]

    # 超管可导任意区域（含 admin 角色）
    batch_id2 = client.post("/api/admin/users/import", headers=auth_of(super_admin), files={
        "file": ("users.xlsx", xlsx_bytes([("33108200078", "超管导", "331082", "", "admin")]),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }).json()["data"]["batch_id"]
    d2 = wait_batch(client, auth_of(super_admin), batch_id2)
    assert d2["success"] == 1 and d2["fail"] == 0


# ---------- 坏文件 / 未知批次 / 空表 ----------

def test_import_bad_file_and_unknown_batch(client, db, import_env):
    seed_regions(db)
    admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    h = auth_of(admin)

    # 非 xlsx 扩展名 → 400
    assert client.post("/api/admin/users/import", headers=h,
                       files={"file": ("users.txt", b"hello", "text/plain")}).status_code == 400
    # xlsx 扩展名但内容非法 → 400
    assert client.post("/api/admin/users/import", headers=h,
                       files={"file": ("bad.xlsx", b"not-an-xlsx",
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code == 400
    # 未知批次 → 404
    assert client.get("/api/admin/users/import/999", headers=h).status_code == 404


def test_import_empty_sheet(client, db, import_env):
    seed_regions(db)
    admin = make_user(db, phone="33100400001", name="临海管理员", role="admin", region="331004")
    batch_id = client.post("/api/admin/users/import", headers=auth_of(admin), files={
        "file": ("empty.xlsx", xlsx_bytes([]),
                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }).json()["data"]["batch_id"]
    d = wait_batch(client, auth_of(admin), batch_id)
    assert d["status"] == "completed" and d["total"] == 0
    assert d["success"] == 0 and d["fail"] == 0 and d["detail"] == []