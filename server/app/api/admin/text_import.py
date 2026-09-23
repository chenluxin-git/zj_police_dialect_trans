"""T17 管理端文本导入：txt 按行 / docx 按段落，句号切分去空白去重；ImportTask 台账轮询 + 撤销（被引用 409 拒撤）
机制见规格 §6.9/§5.1。侧车台账：import_tasks 表仅 {status,error_message,created_at}（规格刻意极简，仅轮询用），
批次明细（文件名/类别/区域/文本 ids/计数）落盘 data/text_imports/{task_id}.json，台账列表/详情/撤销读侧车。
后台执行沿用 BackgroundTasks（TestClient 下响应返回即任务完成，天然可测），会话工厂 _session_factory 供测试重定向。
"""
import json
import os
import re
from datetime import datetime

from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session

from ...core.config import settings
from ...core.database import SessionLocal, get_db
from ...models import Dialect, ImportTask, Recording, Text, User
from ...schemas import ok
from ...services import audit
from ..deps import require_admin, resolve_scope
from .audio_upload import _resolve_region  # 区县级强制裁定（2026-09-22）：三入库口共一定义点

router = APIRouter(prefix="/texts", tags=["管理端-文本导入"])
manage_router = APIRouter(prefix="/text-import-manage", tags=["管理端-文本导入台账"])

_session_factory = SessionLocal  # 后台任务会话工厂（请求级依赖覆盖不可用于后台，测试重定向至此）

_SAMPLE = ["警察同志，请你说明一下案发时的具体情况。", "请出示您的身份证件，配合我们登记。",
           "请不要在公共场所大声喧哗。"]


def _manifest_dir() -> str:
    return settings.text_import_dir  # 容器内指向 /data 卷（.env.docker），防侧车台账落临时层


def _visible_or_403(db: Session, admin: User, task: ImportTask) -> dict:
    """台账辖区隔离（终审 Important#1）：轮询/详情/撤销仅限本辖区批次；
    超管/省管 scope=None 全量可见，市/县管按 manifest region_code ∈ resolve_scope。"""
    m = _load_manifest(task.id)
    scope = resolve_scope(db, admin)
    if scope is not None and m.get("region_code", "") not in scope:
        raise HTTPException(403, "该导入批次不在你的辖区")
    return m


def _manifest_path(task_id: int) -> str:
    return os.path.join(_manifest_dir(), f"{task_id}.json")


def _save_manifest(task_id: int, data: dict) -> None:
    d = _manifest_dir()
    os.makedirs(d, exist_ok=True)
    with open(_manifest_path(task_id), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False)


def _load_manifest(task_id: int) -> dict:
    try:
        with open(_manifest_path(task_id), "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def _split_sentences(line: str) -> list[str]:
    """句号切分（中英文句号），去空白；空白段丢弃"""
    return [p.strip() for p in re.split(r"(?<=[。.])", line) if p.strip()]


def _parse(content: bytes, ext: str) -> list[str]:
    """txt 按行 / docx 按段落 → 句号切分 → 去空白 → 去重（保序）"""
    sentences: list[str] = []
    if ext == ".docx":
        import docx
        from io import BytesIO
        doc = docx.Document(BytesIO(content))
        for para in doc.paragraphs:
            sentences += _split_sentences(para.text)
    else:  # txt
        text = content.decode("utf-8", errors="ignore")
        for line in text.splitlines():
            sentences += _split_sentences(line)
    seen: set[str] = set()
    out: list[str] = []
    for s in sentences:
        if s not in seen:
            seen.add(s)
            out.append(s)
    return out


def _dialect_of(db: Session, region_code: str) -> tuple[str, str]:
    d = db.query(Dialect).filter(Dialect.region_code == region_code).first()
    return (d.name, d.code) if d else ("未指定", "")


def _process_import(task_id: int, content: bytes, ext: str, category: str,
                    region_code: str, file_name: str) -> None:
    db = _session_factory()
    try:
        task = db.get(ImportTask, task_id)
        if task is None:
            return
        task.status = "processing"
        db.commit()
        sentences = _parse(content, ext)
        dialect_name, dialect_code = _dialect_of(db, region_code)
        text_ids: list[int] = []
        for s in sentences:
            t = Text(content=s, dialect=dialect_name, category=category,
                     region_code=region_code, dialect_code=dialect_code)
            db.add(t)
            db.flush()
            text_ids.append(t.id)
        _save_manifest(task_id, {"kind": "text", "file_name": file_name, "category": category,
                                 "region_code": region_code, "total_count": len(text_ids),
                                 "text_ids": text_ids})
        task.status = "completed"
        db.commit()
    except Exception as e:  # noqa: BLE001
        db.rollback()
        task = db.get(ImportTask, task_id)
        if task is not None:
            task.status = "failed"
            task.error_message = str(e)[:500]
        db.commit()
    finally:
        db.close()


@router.get("/template/{fmt}")
def download_template(fmt: str, admin: User = Depends(require_admin)):
    if fmt == "txt":
        body = "\n".join(_SAMPLE).encode("utf-8")
        return Response(content=body, media_type="text/plain",
                        headers={"Content-Disposition": "attachment; filename=texts_template.txt"})
    if fmt == "docx":
        import docx
        from io import BytesIO
        doc = docx.Document()
        for s in _SAMPLE:
            doc.add_paragraph(s)
        buf = BytesIO()
        doc.save(buf)
        return Response(content=buf.getvalue(),
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers={"Content-Disposition": "attachment; filename=texts_template.docx"})
    raise HTTPException(400, "不支持的模板格式")


@router.post("/import")
async def import_texts(
    background: BackgroundTasks,
    request: Request,
    file: UploadFile = File(...),
    category: str = Form(...),
    region_code: str | None = Form(None),
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    file_name = file.filename or "text.txt"
    ext = os.path.splitext(file_name)[1].lower()
    if ext not in (".txt", ".docx"):
        raise HTTPException(400, "仅支持 txt/docx 文件")
    region = _resolve_region(db, admin, region_code)
    task = ImportTask(status="pending")
    db.add(task)
    db.commit()
    _save_manifest(task.id, {"kind": "text", "file_name": file_name, "category": category,
                             "region_code": region, "total_count": 0, "text_ids": []})
    content = await file.read()
    background.add_task(_process_import, task.id, content, ext, category, region, file_name)
    audit.queue_audit(operate_type=audit.OP_CREATE, operate_name="文本导入", user=admin, request=request,
                      operate_condition=(f"执行了[文本导入]功能，操作参数为[文件：{file_name}"
                                         f"||类别：{category}||归属区域：{region}]。"),
                      display=f"导入任务ID={task.id}", data_level=1)
    return ok({"task_id": task.id})


@router.get("/import/{task_id}")
def poll_import(task_id: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    task = db.get(ImportTask, task_id)
    if task is None:
        raise HTTPException(404, "任务不存在")
    m = _visible_or_403(db, admin, task)
    return ok({"task_id": task.id, "status": task.status, "error_message": task.error_message,
               "total_count": m.get("total_count", 0)})


# ---------- 台账（/text-import-manage） ----------

@manage_router.get("")
def list_manage(page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
                admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    scope = resolve_scope(db, admin)
    rows = []
    for t in db.query(ImportTask).order_by(ImportTask.id.desc()).all():
        m = _load_manifest(t.id)
        if m.get("kind") != "text":
            continue
        if scope is not None and m.get("region_code", "") not in scope:  # 台账辖区隔离
            continue
        rows.append({"id": t.id, "status": t.status, "file_name": m.get("file_name", ""),
                     "category": m.get("category", ""), "region_code": m.get("region_code", ""),
                     "total_count": m.get("total_count", 0), "error_message": t.error_message,
                     "created_at": t.created_at.isoformat()})
    total = len(rows)
    start = (page - 1) * page_size
    return ok({"total": total, "page": page, "page_size": page_size,
               "items": rows[start:start + page_size]})


@manage_router.get("/{task_id}")
def manage_detail(task_id: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    task = db.get(ImportTask, task_id)
    if task is None:
        raise HTTPException(404, "任务不存在")
    m = _visible_or_403(db, admin, task)
    sample = [t.content for tid in m.get("text_ids", [])[:10] if (t := db.get(Text, tid))]
    return ok({"id": task.id, "status": task.status, "file_name": m.get("file_name", ""),
               "category": m.get("category", ""), "region_code": m.get("region_code", ""),
               "total_count": m.get("total_count", 0), "error_message": task.error_message,
               "created_at": task.created_at.isoformat(), "sample_texts": sample})


@manage_router.delete("/{task_id}")
def manage_undo(task_id: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    task = db.get(ImportTask, task_id)
    if task is None:
        raise HTTPException(404, "任务不存在")
    m = _visible_or_403(db, admin, task)
    ids = m.get("text_ids", [])
    if not ids:
        db.delete(task)
        db.commit()
        return ok({"deleted": 0})
    if db.query(Recording.text_id).filter(Recording.text_id.in_(ids)).first():
        raise HTTPException(409, "本批次文本已被录音引用，无法撤销")
    db.query(Text).filter(Text.id.in_(ids)).delete(synchronize_session=False)
    db.delete(task)
    db.commit()
    try:
        os.remove(_manifest_path(task_id))
    except OSError:
        pass
    return ok({"deleted": len(ids)})