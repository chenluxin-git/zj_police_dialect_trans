"""T15 批量导入开户：xlsx 模板 + 后台线程逐行校验 + UserImportBatch 台账轮询
机制（Spec §6.6 / §9 后台任务 + 台账轮询）：
- POST /import 解析 xlsx → 建 UserImportBatch → threading.Thread(daemon=True) 逐行处理
- 逐行校验：姓名非空 / 手机号 ^\\d{11}$ 且库内与本批次查重 / 区域码存在且 ∈ 管理范围 / 角色 ∈ user,admin
- 通过行入库：初始密码 = 手机号后 6 位，import_batch_id = 批次号
- 台账 detail 存 JSON 文本（[{phone, ok, msg}]）；status 由 detail 完整度派生
  （表结构无 status 列——Spec §5.1 user_import_batches 字段表，集成裁定）
- 线程会话用模块级 SessionLocal（请求级依赖覆盖不可用于线程，测试重定向至此）
"""
import io
import json
import logging
import re
import threading

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ...core.database import SessionLocal, get_db
from ...core.security import hash_password
from ...models import Region, User, UserImportBatch
from ...schemas import ok
from ...services import audit
from ..deps import require_admin, resolve_scope

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/users", tags=["管理端-批量导入开户"])

PHONE_RE = re.compile(r"^\d{11}$")
IMPORT_ROLES = {"user", "admin"}          # super_admin 不可经导入创建
TEMPLATE_HEADERS = ["手机号", "姓名", "区域码", "单位", "角色"]


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse_rows(content: bytes) -> list[dict]:
    """解析 xlsx 数据行（跳过表头与全空行），返回 [{phone, real_name, region_code, police_station, role}]"""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            cells = [_cell(v) for v in (list(row) + ["", "", "", "", ""])[:5]]
            if not any(cells):
                continue
            rows.append({"phone": cells[0], "real_name": cells[1], "region_code": cells[2],
                         "police_station": cells[3], "role": cells[4]})
        return rows
    finally:
        wb.close()


@router.get("/import-template")
def download_template(admin: User = Depends(require_admin)):
    wb = Workbook()
    ws = wb.active
    ws.title = "批量开户"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["33100400001", "张三", "331004", "临海市公安局××派出所", "user"])
    return _xlsx_response(wb, "users_import_template.xlsx")


@router.post("/import")
async def import_users(request: Request, file: UploadFile = File(...),
                       admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "仅支持 .xlsx 文件")
    content = await file.read()
    try:
        rows = _parse_rows(content)
    except Exception:
        logger.exception("开户文件解析失败 %s", file.filename)
        raise HTTPException(400, "文件解析失败，请使用下载的模板填写")

    batch = UserImportBatch(file_name=file.filename or "", total=len(rows), success=0, fail=0,
                            detail="[]", created_by=admin.id)
    db.add(batch)
    db.commit()
    scope = resolve_scope(db, admin)   # 请求会话内先取好，线程里不再依赖请求对象
    threading.Thread(target=_process_batch, args=(batch.id, rows, scope),
                     daemon=True, name=f"import-users-{batch.id}").start()
    audit.queue_audit(operate_type=audit.OP_CREATE, operate_name="用户批量导入", user=admin, request=request,
                      operate_condition=(f"执行了[用户批量开户导入]功能，操作参数为[文件：{file.filename}"
                                         f"||总行数：{len(rows)}]。"),
                      display=f"导入批次ID={batch.id}", data_level=2)
    return ok({"batch_id": batch.id})


def _save_progress(db: Session, batch_id: int, detail: list[dict], success: int, fail: int) -> None:
    batch = db.get(UserImportBatch, batch_id)
    if batch is None:
        return
    batch.detail = json.dumps(detail, ensure_ascii=False)
    batch.success = success
    batch.fail = fail
    db.commit()


def _process_batch(batch_id: int, rows: list[dict], scope: list[str] | None) -> None:
    """后台线程：逐行校验 + 写库 + 更新台账（detail 逐行追加）"""
    db = SessionLocal()
    detail: list[dict] = []
    success = fail = 0
    try:
        seen: set[str] = set()
        region_codes = set(db.scalars(select(Region.code)).all())
        for row in rows:
            phone, name, region = row["phone"], row["real_name"], row["region_code"]
            role, station = row["role"], row["police_station"]
            if not name:
                msg = "姓名不能为空"
            elif not PHONE_RE.fullmatch(phone):
                msg = "手机号必须为11位数字"
            elif phone in seen or db.query(User).filter(User.phone == phone).first() is not None:
                msg = "手机号已存在"
            elif region not in region_codes:
                msg = f"区域码不存在：{region}"
            elif scope is not None and region not in scope:
                msg = "区域超出管理范围"
            elif role not in IMPORT_ROLES:
                msg = "角色不合法（仅支持 user/admin）"
            else:
                db.add(User(phone=phone, password_hash=hash_password(phone[-6:]),
                            real_name=name, police_station=station,
                            region_code=region, role=role, import_batch_id=batch_id))
                db.commit()
                seen.add(phone)
                success += 1
                detail.append({"phone": phone, "ok": True, "msg": "ok"})
                _save_progress(db, batch_id, detail, success, fail)
                continue
            fail += 1
            detail.append({"phone": phone, "ok": False, "msg": msg})
            _save_progress(db, batch_id, detail, success, fail)
    except Exception:
        logger.exception("开户批次 %s 处理失败", batch_id)
        db.rollback()
    finally:
        db.close()


@router.get("/import/{batch_id}")
def import_status(batch_id: int, admin: User = Depends(require_admin), db: Session = Depends(get_db)):
    batch = db.get(UserImportBatch, batch_id)
    if batch is None:
        raise HTTPException(404, "批次不存在")
    detail = json.loads(batch.detail or "[]")
    status = "completed" if len(detail) >= batch.total else "processing"
    return ok({"batch_id": batch.id, "file_name": batch.file_name, "status": status,
               "total": batch.total, "success": batch.success, "fail": batch.fail, "detail": detail})
