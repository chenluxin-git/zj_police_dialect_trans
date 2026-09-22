"""T17 文本导入测试：txt/docx 句号切分去空白去重、ImportTask 台账轮询、模板下载、撤销（被引用 409 拒撤）
后台任务说明：实现用 BackgroundTasks（沿用旧项目机制），TestClient 下 POST 返回即任务已完成；
后台任务自带会话工厂（模块级 _session_factory），测试重定向到 db fixture 同一引擎（db.get_bind()）。
"""
import io

import pytest
from sqlalchemy.orm import sessionmaker

from app.core.security import create_token
from app.models import Dialect, ImportTask, Recording, Region, Text, User  # 模块级导入：先注册模型
from tests.conftest import make_user


def seed_regions_and_dialects(db):
    db.add_all([
        Region(code="330000", name="浙江省", level="province", parent_code=None),
        Region(code="331000", name="台州市", level="city", parent_code="330000"),
        Region(code="331004", name="路桥区", level="district", parent_code="331000"),
    ])
    db.add(Dialect(code="dh_lq", name="路桥方言", region_code="331004"))
    db.commit()


def auth_of(user):
    return {"Authorization": f"Bearer {create_token(str(user.id))}"}


@pytest.fixture()
def import_env(db, tmp_path, monkeypatch):
    """后台任务会话 → db 同引擎；侧车台账落盘 → tmp_path（不污染 server/data）"""
    from app.api.admin import text_import as ti
    monkeypatch.setattr(ti, "_session_factory",
                        sessionmaker(bind=db.get_bind(), expire_on_commit=False))
    monkeypatch.setattr(ti, "_manifest_dir", lambda: str(tmp_path))
    return ti


def _post_import(client, h, content, category="police", filename="t.txt", mime="text/plain"):
    return client.post("/api/admin/texts/import", headers=h,
                       data={"category": category},
                       files={"file": (filename, content if isinstance(content, bytes) else content.encode("utf-8"), mime)})


# ---------- txt 导入 + 轮询 ----------

def test_txt_import_and_poll(client, db, import_env):
    seed_regions_and_dialects(db)
    admin = make_user(db, phone="33100400001", name="路桥管理员", role="admin", region="331004")
    h = auth_of(admin)

    r = _post_import(client, h, "第一句。第二句。第三句没有句号")
    assert r.status_code == 200 and r.json()["code"] == 0
    task_id = r.json()["data"]["task_id"]

    t = client.get(f"/api/admin/texts/import/{task_id}", headers=h)
    assert t.json()["data"]["status"] == "completed"
    assert t.json()["data"]["total_count"] == 3  # 句号切分：末句无句号亦按一行一句保留

    texts = db.query(Text).all()
    assert {x.content for x in texts} == {"第一句。", "第二句。", "第三句没有句号"}
    assert all(x.category == "police" and x.dialect == "路桥方言"
               and x.dialect_code == "dh_lq" and x.region_code == "331004" for x in texts)


def test_txt_import_dedup(client, db, import_env):
    seed_regions_and_dialects(db)
    admin = make_user(db, phone="33100400001", name="路桥管理员", role="admin", region="331004")
    _post_import(client, auth_of(admin), "重复句。重复句。独特句。", category="life")
    assert db.query(Text).count() == 2  # 去重


# ---------- docx 导入 ----------

def test_docx_import(client, db, import_env):
    import docx
    seed_regions_and_dialects(db)
    admin = make_user(db, phone="33100400001", name="路桥管理员", role="admin", region="331004")

    d = docx.Document()
    d.add_paragraph("甲句。乙句。")
    d.add_paragraph("丙句。")
    buf = io.BytesIO()
    d.save(buf)

    r = _post_import(client, auth_of(admin), buf.getvalue(), filename="t.docx",
                     mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert r.status_code == 200
    assert {x.content for x in db.query(Text).all()} == {"甲句。", "乙句。", "丙句。"}


# ---------- 模板下载 ----------

def test_template_download(client, db):
    seed_regions_and_dialects(db)
    admin = make_user(db, phone="33100400001", name="路桥管理员", role="admin", region="331004")
    h = auth_of(admin)

    r = client.get("/api/admin/texts/template/txt", headers=h)
    assert r.status_code == 200 and "。" in r.text
    r2 = client.get("/api/admin/texts/template/docx", headers=h)
    assert r2.status_code == 200 and r2.content[:2] == b"PK"  # docx 为 zip 容器


# ---------- 台账列表 / 详情 / 撤销 ----------

def test_undo_and_409(client, db, import_env):
    seed_regions_and_dialects(db)
    admin = make_user(db, phone="33100400001", name="路桥管理员", role="admin", region="331004")
    h = auth_of(admin)

    _post_import(client, h, "一句。二句。")
    task_id = db.query(ImportTask).order_by(ImportTask.id.desc()).first().id
    assert db.query(Text).count() == 2

    # 台账列表 + 详情
    assert client.get("/api/admin/text-import-manage", headers=h).json()["data"]["total"] == 1
    detail = client.get(f"/api/admin/text-import-manage/{task_id}", headers=h).json()["data"]
    assert detail["total_count"] == 2 and detail["status"] == "completed"

    # 撤销 → 文本删除 + 台账清空
    r = client.delete(f"/api/admin/text-import-manage/{task_id}", headers=h)
    assert r.json()["data"]["deleted"] == 2
    assert db.query(Text).count() == 0
    assert client.get("/api/admin/text-import-manage", headers=h).json()["data"]["total"] == 0

    # 再导入一批，其中一条被录音引用 → 409 拒撤
    _post_import(client, h, "甲。乙。")
    task2 = db.query(ImportTask).order_by(ImportTask.id.desc()).first().id
    t_objs = db.query(Text).all()
    officer = make_user(db, phone="33100400002", name="路桥民警", role="user", region="331004")
    db.add(Recording(user_id=officer.id, text_id=t_objs[0].id, file_path="", file_size=0,
                     duration=0.0, region_code="331004", dialect_code="dh_lq", qc_status="pending"))
    db.commit()
    assert client.delete(f"/api/admin/text-import-manage/{task2}", headers=h).status_code == 409
    assert db.query(Text).count() == 2  # 拒撤，文本保留