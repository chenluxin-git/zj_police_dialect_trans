"""T14 管理端用户管理测试：scope 列表/筛选/统计列、创建（初始密码后6位+登录）、修改、删除（有录音禁删）、Excel 导出
口径：计划 T14 Step 1（县管只见本县用户、县管不可改 super_admin 403、删除有录音用户 400、
导出 xlsx content-type 含 spreadsheet 且 bytes 非空）+ 筛选/进度列/创建权限补充。
多用户令牌：conftest auth_header 硬编码 id=1，本文件一律 create_token(str(user.id)) 自造（分发手册公共事实）。
"""
import io

from openpyxl import load_workbook

from app.core.security import create_token
from app.models import (Annotation, AudioFile, Recording, Region, Task,  # 模块级导入：conftest 建表前须已注册模型
                        Text, User)
from tests.conftest import make_user


# ---------- 造数 ----------

def seed_regions(db):
    db.add_all([
        Region(code="330000", name="浙江省", level="province", parent_code=None),
        Region(code="331000", name="台州市", level="city", parent_code="330000"),
        Region(code="331004", name="临海市", level="district", parent_code="331000"),
        Region(code="331082", name="三门县", level="district", parent_code="331000"),
        Region(code="330100", name="杭州市", level="city", parent_code="330000"),
        Region(code="330105", name="拱墅区", level="district", parent_code="330100"),
    ])
    db.commit()


def make_admin(db, phone, name, role, region):
    return make_user(db, phone=phone, name=name, role=role, region=region)


def auth_of(user):
    return {"Authorization": f"Bearer {create_token(str(user.id))}"}


# ---------- 列表：scope / 筛选 / 分页 ----------

def test_list_scope_filters_pagination(client, db):
    seed_regions(db)
    u_lh = make_user(db, phone="33100400002", name="临海民警")                # role=user，辖区 331004
    make_user(db, phone="33108200002", name="三门民警", region="331082")
    make_user(db, phone="33010500002", name="杭州民警", region="330105",
              station="杭州市公安局××派出所")
    tz_admin = make_admin(db, "33100000001", "台州管理员", "admin", "331000")
    county_admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")
    super_admin = make_admin(db, "33000000001", "省超管", "super_admin", "330000")

    # 民警访问管理端 → 403
    assert client.get("/api/admin/users", headers=auth_of(u_lh)).status_code == 403

    # 县管：仅本县（331004）→ 临海民警 + 本人
    r = client.get("/api/admin/users", headers=auth_of(county_admin))
    assert r.status_code == 200 and r.json()["code"] == 0
    data = r.json()["data"]
    assert data["total"] == 2
    assert {i["real_name"] for i in data["items"]} == {"临海民警", "临海管理员"}

    # 市管：台州辖区（331000/331004/331082）→ 4；超管：全部 6
    assert client.get("/api/admin/users", headers=auth_of(tz_admin)).json()["data"]["total"] == 4
    assert client.get("/api/admin/users", headers=auth_of(super_admin)).json()["data"]["total"] == 6

    # 筛选：real_name / phone / role / station
    assert client.get("/api/admin/users", headers=auth_of(super_admin),
                      params={"real_name": "三门"}).json()["data"]["total"] == 1
    assert client.get("/api/admin/users", headers=auth_of(super_admin),
                      params={"phone": "330105"}).json()["data"]["total"] == 1
    assert client.get("/api/admin/users", headers=auth_of(super_admin),
                      params={"role": "admin"}).json()["data"]["total"] == 2
    assert client.get("/api/admin/users", headers=auth_of(super_admin),
                      params={"station": "临海市公安局××派出所"}).json()["data"]["total"] == 5

    # region_code 筛选（展开下级 ∩ scope）：超管按 331000 → 台州 4；县管按 331000 → ∩ scope 仍本县 2
    assert client.get("/api/admin/users", headers=auth_of(super_admin),
                      params={"region_code": "331000"}).json()["data"]["total"] == 4
    assert client.get("/api/admin/users", headers=auth_of(county_admin),
                      params={"region_code": "331000"}).json()["data"]["total"] == 2

    # 分页
    paged = client.get("/api/admin/users", headers=auth_of(super_admin),
                       params={"page": 1, "page_size": 2}).json()["data"]
    assert len(paged["items"]) == 2 and paged["total"] == 6
    assert len(client.get("/api/admin/users", headers=auth_of(super_admin),
                          params={"page": 4, "page_size": 2}).json()["data"]["items"]) == 0


# ---------- 列表：统计列（progress_map 口径 + 任务进度） ----------

def test_list_counts_and_task_progress(client, db):
    seed_regions(db)
    u = make_user(db, phone="33100400002", name="临海民警")
    admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")
    h = auth_of(admin)

    # recordings 表有 UNIQUE(user_id, text_id)，一人一条文本只录一次 → 需三条不同文本
    for i, status in enumerate(["passed", "passed", "pending"]):
        t = Text(content=f"你们不要吵了第{i}句", dialect="临海方言", category="police",
                 region_code="331004", dialect_code="dh_lh")
        db.add(t)
        db.flush()
        db.add(Recording(user_id=u.id, text_id=t.id, file_path=f"r{i}.wav", file_size=10,
                         duration=3.0, region_code="331004", dialect_code="dh_lh", qc_status=status))
    db.commit()
    af = AudioFile(file_path="a.mp3", file_name="a.mp3", duration=5.0,
                   region_code="331004", dialect_code="dh_lh")
    db.add(af)
    db.commit()
    db.add(Annotation(file_id=af.id, annotator_id=u.id, is_dialect=True,
                      translation="下雨了，衣服要收起来了", region_code="331004"))
    db.add(Task(user_id=u.id, type="recording", target_count=10, base_count=1,
                status="active", note="首批", created_by=admin.id))
    db.commit()

    data = client.get("/api/admin/users", headers=h).json()["data"]
    item = next(i for i in data["items"] if i["id"] == u.id)
    assert item["recording_count"] == 2        # 仅 qc_status=passed（progress_map 口径）
    assert item["annotation_count"] == 1
    tp = item["task_progress"]
    assert tp["type"] == "recording" and tp["target_count"] == 10
    assert tp["base_count"] == 1 and tp["done"] == 1 and tp["status"] == "active"

    # 管理员本人无 active 任务 → None
    a_item = next(i for i in data["items"] if i["id"] == admin.id)
    assert a_item["task_progress"] is None


# ---------- 创建：初始密码 = 手机号后 6 位 ----------

def test_create_user_initial_password_and_login(client, db):
    seed_regions(db)
    super_admin = make_admin(db, "33000000001", "省超管", "super_admin", "330000")
    h = auth_of(super_admin)

    r = client.post("/api/admin/users", headers=h, json={
        "phone": "13900000001", "real_name": "新民警", "region_code": "331004",
        "police_station": "临海市公安局", "role": "user"})
    assert r.status_code == 200 and r.json()["code"] == 0
    created = r.json()["data"]
    assert created["role"] == "user" and created["phone"] == "13900000001"

    # 初始密码 = 手机号后 6 位 → 可登录
    login = client.post("/api/auth/login", json={"phone": "13900000001", "password": "000001"})
    assert login.status_code == 200 and login.json()["data"]["token"]

    # 重复手机号 / 手机号格式 / 区域不存在 / 角色非法 → 400
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "13900000001", "real_name": "重复", "region_code": "331004",
        "role": "user"}).status_code == 400
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "123", "real_name": "短号", "region_code": "331004",
        "role": "user"}).status_code == 400
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "13900000002", "real_name": "坏区域", "region_code": "999999",
        "role": "user"}).status_code == 400
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "13900000003", "real_name": "坏角色", "region_code": "331004",
        "role": "boss"}).status_code == 400


def test_create_user_permissions(client, db):
    seed_regions(db)
    county_admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")
    h = auth_of(county_admin)

    # 辖区外区域 → 403
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "13900000002", "real_name": "越界账号", "region_code": "331082",
        "role": "user"}).status_code == 403
    # admin 创建 super_admin → 403
    assert client.post("/api/admin/users", headers=h, json={
        "phone": "13900000003", "real_name": "越权超管", "region_code": "331004",
        "role": "super_admin"}).status_code == 403
    # 辖区内创建 admin（非 super_admin）→ 200
    r = client.post("/api/admin/users", headers=h, json={
        "phone": "13900000004", "real_name": "辖区管理员", "region_code": "331004",
        "role": "admin"})
    assert r.status_code == 200 and r.json()["data"]["role"] == "admin"


# ---------- 修改 ----------

def test_update_user(client, db):
    seed_regions(db)
    u = make_user(db, phone="33100400002", name="临海民警")
    hz_user = make_user(db, phone="33010500002", name="杭州民警", region="330105")
    super_admin = make_admin(db, "33000000001", "省超管", "super_admin", "330000")
    county_admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")

    # 县管改本辖区用户：姓名/单位/角色/重置密码 → 登录验证新密码
    r = client.put(f"/api/admin/users/{u.id}", headers=auth_of(county_admin), json={
        "real_name": "改名民警", "police_station": "临海市公安局杜桥派出所",
        "role": "admin", "password": "999999"})
    assert r.status_code == 200 and r.json()["data"]["real_name"] == "改名民警"
    login = client.post("/api/auth/login", json={"phone": "33100400002", "password": "999999"})
    assert login.status_code == 200

    # 县管改 super_admin / 改辖区外用户 / 移到辖区外区域 / 升为 super_admin → 403
    assert client.put(f"/api/admin/users/{super_admin.id}", headers=auth_of(county_admin),
                      json={"real_name": "x"}).status_code == 403
    assert client.put(f"/api/admin/users/{hz_user.id}", headers=auth_of(county_admin),
                      json={"real_name": "x"}).status_code == 403
    assert client.put(f"/api/admin/users/{u.id}", headers=auth_of(county_admin),
                      json={"region_code": "331082"}).status_code == 403
    assert client.put(f"/api/admin/users/{u.id}", headers=auth_of(county_admin),
                      json={"role": "super_admin"}).status_code == 403

    # 超管不受限：可改任意用户
    r2 = client.put(f"/api/admin/users/{hz_user.id}", headers=auth_of(super_admin),
                    json={"real_name": "杭州改名"})
    assert r2.status_code == 200 and r2.json()["data"]["real_name"] == "杭州改名"

    # 未知用户 → 404
    assert client.put("/api/admin/users/999", headers=auth_of(county_admin),
                      json={"real_name": "x"}).status_code == 404


# ---------- 删除 ----------

def test_delete_user(client, db):
    seed_regions(db)
    u = make_user(db, phone="33100400002", name="临海民警")
    county_admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")
    super_admin = make_admin(db, "33000000001", "省超管", "super_admin", "330000")
    h = auth_of(county_admin)

    # 有录音 → 400
    t = Text(content="请出示健康码", dialect="杭州方言", category="police",
             region_code="331004", dialect_code="dh_lh")
    db.add(t)
    db.commit()
    db.add(Recording(user_id=u.id, text_id=t.id, file_path="r.wav", file_size=1,
                     duration=1.0, region_code="331004", dialect_code="dh_lh",
                     qc_status="pending"))
    db.commit()
    assert client.delete(f"/api/admin/users/{u.id}", headers=h).status_code == 400

    # 无录音 → 200 且列表不再包含
    db.query(Recording).delete()
    db.commit()
    assert client.delete(f"/api/admin/users/{u.id}", headers=h).status_code == 200
    items = client.get("/api/admin/users", headers=h).json()["data"]["items"]
    assert all(i["id"] != u.id for i in items)

    # 县管删 super_admin → 403；未知用户 → 404
    assert client.delete(f"/api/admin/users/{super_admin.id}", headers=h).status_code == 403
    assert client.delete("/api/admin/users/999", headers=h).status_code == 404


# ---------- Excel 导出 ----------

def test_export_xlsx(client, db):
    seed_regions(db)
    make_user(db, phone="33100400002", name="临海民警")
    make_user(db, phone="33108200002", name="三门民警", region="331082")
    county_admin = make_admin(db, "33100400001", "临海管理员", "admin", "331004")
    super_admin = make_admin(db, "33000000001", "省超管", "super_admin", "330000")

    r = client.get("/api/admin/users/export", headers=auth_of(super_admin))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert len(r.content) > 0
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["手机号", "姓名", "角色", "区域", "单位", "录音数", "标注数"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 4                       # 2 民警 + 县管 + 超管
    assert "33100400002" in {row[0] for row in rows}

    # 县管导出仅本辖区
    r2 = client.get("/api/admin/users/export", headers=auth_of(county_admin))
    rows2 = list(load_workbook(io.BytesIO(r2.content)).active.iter_rows(min_row=2, values_only=True))
    assert len(rows2) == 2
    assert all(row[3] == "331004" for row in rows2)
